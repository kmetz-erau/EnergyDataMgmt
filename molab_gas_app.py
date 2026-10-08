# /// script
# requires-python = ">=3.10"
# dependencies = [
#     "marimo",
#     "openpyxl",
#     "requests",
# ]
# ///

"""MoLab/marimo app for combining gas bills with utility totals."""

import marimo

app = marimo.App(width="full")


@app.cell
def _():
    import csv
    import io
    import re
    from calendar import month_name
    from datetime import date

    import marimo as mo
    import openpyxl
    import requests

    return csv, date, io, mo, month_name, openpyxl, re, requests


@app.cell
def _(mo):
    gas_bills_upload = mo.ui.file(
        filetypes=[".csv"],
        multiple=False,
        kind="area",
        label="Upload gas_bills.csv",
    )

    utilities_upload = mo.ui.file(
        filetypes=[".xlsx", ".xlsm"],
        multiple=False,
        kind="area",
        label="Upload UTILITIES TOTALS FY 2026-2027.xlsx",
    )

    dedupe = mo.ui.checkbox(
        value=True,
        label="Skip duplicate provider/account/month rows",
    )

    run_button = mo.ui.run_button(
        label="Combine gas data and apply mapping",
        kind="success",
    )

    mo.vstack(
        [
            mo.md("## Gas Bills Combiner"),
            mo.md(
                "Upload the gas bills CSV and the utilities totals workbook. "
                "The gas account mapping is loaded from GitHub."
            ),
            gas_bills_upload,
            utilities_upload,
            dedupe,
            run_button,
        ],
        gap=1,
    )

    return dedupe, gas_bills_upload, run_button, utilities_upload


@app.cell
def _(csv, date, io, month_name, openpyxl, re):
    MAPPING_URL = (
        "https://raw.githubusercontent.com/kmetz-erau/EnergyDataMgmt/main/"
        "gas_account_mapping.csv"
    )

    def clean(value):
        return "" if value is None else str(value).strip()

    def account_key(value):
        value = clean(value)
        if not value:
            return ""
        return value.lstrip("0") or "0"

    def provider_key(value):
        return re.sub(r"\s+", " ", clean(value)).upper()

    def parse_csv(data):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        if not reader.fieldnames:
            raise ValueError("gas_bills.csv has no header row.")
        return list(reader.fieldnames), list(reader)

    def csv_bytes(fields, rows):
        output = io.StringIO(newline="")
        writer = csv.DictWriter(
            output,
            fieldnames=fields,
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(
            {field: row.get(field, "") for field in fields}
            for row in rows
        )
        return output.getvalue().encode("utf-8")

    def month_number(value):
        text = clean(value).upper()
        for number in range(1, 13):
            if text in {month_name[number].upper(), month_name[number][:3].upper()}:
                return number
        return None

    def fiscal_start_year(workbook):
        sheet = workbook["GAS"]
        header_text = " ".join(
            clean(cell.value)
            for row in sheet.iter_rows(min_row=1, max_row=2)
            for cell in row
            if cell.value is not None
        )
        match = re.search(r"(?:FY\s*)?(\d{2})\s*/\s*(\d{2})", header_text)
        if match:
            return 2000 + int(match.group(1))
        return date.today().year

    def parse_utilities_gas(data):
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )

        if "GAS" not in workbook.sheetnames:
            raise ValueError("The workbook does not contain a GAS sheet.")

        sheet = workbook["GAS"]
        values = list(sheet.iter_rows(values_only=True))

        if len(values) < 4:
            raise ValueError("The GAS sheet does not contain enough rows.")

        # Row 2 contains month labels repeated for THERMS/COST.
        # Row 3 contains the metric labels: THERMS, COST, THERMS, COST...
        month_header = list(values[1])
        metric_header = list(values[2])
        start_year = fiscal_start_year(workbook)
        rows = []

        # GAS sheet columns are: Bldg #, Name, TECO account, ERAU account,
        # followed by repeating THERMS/COST pairs for July through June.
        for source_row in values[3:]:
            if len(source_row) < 4:
                continue

            utility_account = clean(source_row[2])
            if not utility_account or utility_account.lower() in {"none", "nan"}:
                continue

            for column in range(4, len(source_row), 2):
                if column >= len(month_header):
                    break

                month = month_number(month_header[column])
                if not month:
                    continue

                therms = source_row[column]
                cost = source_row[column + 1] if column + 1 < len(source_row) else None

                if therms in (None, "") and cost in (None, ""):
                    continue

                calendar_year = start_year if month >= 7 else start_year + 1
                rows.append(
                    {
                        "provider": "TECO",
                        "utility_account": utility_account,
                        "erau_account": clean(source_row[3]),
                        "account_name": clean(source_row[1]),
                        "building_id": clean(source_row[0]),
                        "billMonth": date(calendar_year, month, 1).isoformat(),
                        "therms": clean(therms),
                        "cost": clean(cost),
                        "source_file": "UTILITIES TOTALS GAS sheet",
                        "gas_key": (
                            f"TECO|{utility_account}|"
                            f"{calendar_year:04d}-{month:02d}"
                        ),
                    }
                )

        return rows

    return (
        MAPPING_URL,
        account_key,
        csv_bytes,
        parse_csv,
        parse_utilities_gas,
        provider_key,
    )


@app.cell
def _(MAPPING_URL, account_key, csv_bytes, parse_csv, parse_utilities_gas, provider_key):
    def combine_gas(gas_data, utilities_data, mapping_data, dedupe):
        gas_fields, gas_rows = parse_csv(gas_data)
        utility_rows = parse_utilities_gas(utilities_data)

        fields = list(gas_fields)
        for field in utility_rows[0].keys() if utility_rows else []:
            if field not in fields:
                fields.append(field)

        identities = set()
        if dedupe:
            for row in gas_rows:
                identity = (
                    provider_key(row.get("provider")),
                    account_key(row.get("utility_account")),
                    str(row.get("billMonth", "")).strip()[:7],
                )
                if all(identity):
                    identities.add(identity)

        appended = 0
        skipped = 0

        for utility_row in utility_rows:
            identity = (
                provider_key(utility_row.get("provider")),
                account_key(utility_row.get("utility_account")),
                str(utility_row.get("billMonth", "")).strip()[:7],
            )

            if dedupe and all(identity) and identity in identities:
                skipped += 1
                continue

            output_row = {
                field: utility_row.get(field, "")
                for field in fields
            }
            gas_rows.append(output_row)
            identities.add(identity)
            appended += 1

        mapping_fields, mapping_rows = parse_csv(mapping_data)
        required_mapping_fields = {"provider", "utility_account"}
        missing = required_mapping_fields.difference(mapping_fields)
        if missing:
            raise ValueError(
                "The GitHub mapping is missing: "
                + ", ".join(sorted(missing))
            )

        mapping_by_account = {}
        for mapping_row in mapping_rows:
            key = (
                provider_key(mapping_row.get("provider")),
                account_key(mapping_row.get("utility_account")),
            )
            if not all(key):
                continue
            mapping_by_account[key] = mapping_row

        output_fields = list(fields)
        added_mapping_fields = []
        for field in mapping_fields:
            if field in {"provider", "utility_account"}:
                continue
            output_field = field if field not in output_fields else f"mapping_{field}"
            added_mapping_fields.append((field, output_field))
            output_fields.append(output_field)

        matched = 0
        unmatched = 0
        output_rows = []

        for row in gas_rows:
            key = (
                provider_key(row.get("provider")),
                account_key(row.get("utility_account")),
            )
            mapping_row = mapping_by_account.get(key)
            if mapping_row:
                matched += 1
            else:
                unmatched += 1

            output_row = {
                field: row.get(field, "")
                for field in fields
            }
            for source_field, output_field in added_mapping_fields:
                output_row[output_field] = (
                    mapping_row.get(source_field, "")
                    if mapping_row
                    else ""
                )
            output_rows.append(output_row)

        stats = {
            "original gas bill rows": len(gas_rows) - appended,
            "utility workbook rows found": len(utility_rows),
            "utility rows appended": appended,
            "duplicates skipped": skipped,
            "mapping matches": matched,
            "mapping unmatched": unmatched,
            "rows written": len(output_rows),
        }

        return output_fields, output_rows, stats

    return combine_gas


@app.cell
def _(
    combine_gas,
    dedupe,
    gas_bills_upload,
    mo,
    requests,
    run_button,
    utilities_upload,
    csv_bytes,
    MAPPING_URL,
):
    if not run_button.value:
        result_display = mo.md(
            "Upload both files, then click **Combine gas data and apply mapping**."
        )
    elif not gas_bills_upload.value:
        result_display = mo.callout(
            mo.md("Upload `gas_bills.csv` first."),
            kind="warn",
        )
    elif not utilities_upload.value:
        result_display = mo.callout(
            mo.md("Upload the utilities totals `.xlsx` file first."),
            kind="warn",
        )
    else:
        try:
            gas_data = gas_bills_upload.value[0].contents
            utilities_data = utilities_upload.value[0].contents

            response = requests.get(MAPPING_URL, timeout=30)
            response.raise_for_status()
            mapping_data = response.content

            fields, rows, stats = combine_gas(
                gas_data,
                utilities_data,
                mapping_data,
                dedupe=dedupe.value,
            )
            result = csv_bytes(fields, rows)

            summary_lines = [
                "### Gas data combined",
                "",
                "| Metric | Count |",
                "|---|---:|",
            ]
            summary_lines.extend(
                f"| {label} | {value} |"
                for label, value in stats.items()
            )

            result_display = mo.vstack(
                [
                    mo.md("\n".join(summary_lines)),
                    mo.download(
                        data=result,
                        filename="gas_bills_combined_mapped.csv",
                        mimetype="text/csv",
                        label="Download gas_bills_combined_mapped.csv",
                    ),
                ]
            )
        except Exception as exc:
            result_display = mo.callout(
                mo.md(f"**Could not process the files:** {exc}"),
                kind="danger",
            )

    result_display

